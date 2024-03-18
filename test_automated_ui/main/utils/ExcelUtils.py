import openpyxl
from openpyxl.styles import PatternFill

def getRowCount(file, sheetname):
    try:
        workbook = openpyxl.load_workbook(file)
        sheet = workbook[sheetname]
        return (sheet.max_row)
    except FileNotFoundError:
        print(f"File '{file}' not found.")
        return None
    except KeyError:
        print(f"Sheet '{sheetname}' not found in '{file}'.")
        return None
    
def getColumnCount(file, sheetname):
    try:
        workbook = openpyxl.load_workbook(file)
        sheet = workbook[sheetname]
        return sheet.max_column
    except FileNotFoundError:
        print(f"File '{file}' not found.")
        return None
    except KeyError:
        print(f"Sheet '{sheetname}' not found in '{file}'.")
        return None
    
def readData(file,sheetname,rownum,columnno):
    try:
        workbook = openpyxl.load_workbook(file)
        sheet = workbook[sheetname]
        return sheet.cell(rownum,columnno).value
    except FileNotFoundError:
        print(f"File '{file}' not found.")
        return None
    except KeyError:
        print(f"Sheet '{sheetname}' not found in '{file}'.")
        return None